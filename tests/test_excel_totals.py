"""Testes dos totais da planilha Excel (FASE 3 - correção).

Verifica que a linha TOTAIS grava a fórmula de matriz COM valor em cache —
sem isso, o LibreOffice/Excel em Modo de Exibição Protegido (arquivo baixado
da internet, MOTW) exibe a célula vazia, pois não recalcula fórmulas ao abrir.
"""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from backend.core.excel_service import count_markers, generate_excel_file

BASE_COLUMNS = [
    'DATA ENTRADA', 'ENTRADA', 'DATA SAÍDA', 'SAÍDA',
    'TRABALHADA', 'HORA JUSTIFICADA', 'STATUS',
    'HT', 'HJ', 'ST', 'ADN',
]


def _year_dataframe() -> pd.DataFrame:
    """Dataframe no formato produzido por generate_dataframe (ano único).

    A coluna EDITAR já foi removida (como no fluxo real), deixando as
    colunas derivadas HT/HJ/ST/ADN nas posições 7..10 (planilha H..K).
    """
    header_message = 'PONTO DIGITAL - FULANO - CPF: 123.456.789-01 - JANEIRO/2024'
    rows = [
        [header_message] + [''] * (len(BASE_COLUMNS) - 1),
        list(BASE_COLUMNS),
        ['01/01/2024', '08:00:00', '01/01/2024', '17:00:00', '8:00:00', '', 'APROVADO', 1, '', 1, ''],
        ['02/01/2024', '08:00:00', '02/01/2024', '17:00:00', '8:00:00', '', 'APROVADO', 1, '', 1, ''],
        ['03/01/2024', '08:00:00', '03/01/2024', '17:00:00', '8:00:00', '', 'REPROVADO', '', '', '', ''],
        ['04/01/2024', '', '04/01/2024', '', '---', '---', '---', '', '', '', 1],
    ]
    totals = ['TOTAIS'] + [''] * (len(BASE_COLUMNS) - 1)
    df = pd.DataFrame(rows + [totals], columns=BASE_COLUMNS, dtype=object)
    # HT=2, HJ=0, ST=2, ADN=1 (a última linha marca apenas ADN, como em férias/ausência)
    return df


def _load(filepath: Path, data_only: bool):
    return load_workbook(filepath, data_only=data_only)


def _totals_row(worksheet):
    for row in worksheet.iter_rows():
        if row[0].value == 'TOTAIS':
            return row
    raise AssertionError('Linha TOTAIS não encontrada')


@pytest.fixture
def generated(tmp_path):
    df = _year_dataframe()
    out = Path(tmp_path) / 'totais.xlsx'
    generate_excel_file(
        data_dic={2024: df},
        employee_name='FULANO',
        cpf='123.456.789-01',
        output_path=out,
        password='x',
    )
    return out


class TestTotaisCached:

    def test_count_markers(self):
        df = _year_dataframe()
        # Intervalo entre o cabeçalho (linha 2) e o TOTAIS (linha 7): linhas 2..6
        assert count_markers(df, 2, 6, 7) == 2   # HT  (coluna H)
        assert count_markers(df, 2, 6, 8) == 0   # HJ  (coluna I)
        assert count_markers(df, 2, 6, 9) == 2   # ST  (coluna J)
        assert count_markers(df, 2, 6, 10) == 1  # ADN (coluna K)

    def test_totais_have_cached_numeric_values(self, generated):
        """data_only=True lê os totais calculados (mesmo sem recalcular)."""
        wb = _load(generated, data_only=True)
        row = _totals_row(wb.active)
        values = [c.value for c in row[7:11]]
        assert values == [2, 0, 2, 1]

    def test_totais_still_are_formulas(self, generated):
        """A fórmula de soma continua presente (paridade com o desktop)."""
        wb = _load(generated, data_only=False)
        row = _totals_row(wb.active)
        for cell in row[7:11]:
            text = getattr(cell.value, 'text', None) or cell.value
            assert str(text).startswith('=SUM(')

    def test_protecao_nao_quebra_download_view(self, generated):
        """Planilha continua protegida, mas com totais visíveis em cache."""
        wb = _load(generated, data_only=True)
        assert wb.active.protection.sheet is True or wb.active.protection.password is not None